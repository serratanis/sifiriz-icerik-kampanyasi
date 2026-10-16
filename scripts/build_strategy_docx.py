from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables" / "SifirIz-Marka-Stratejisi.docx"

FOREST = "173F3A"
SAGE = "A7C4A0"
APRICOT = "E76F51"
CREAM = "F5F0E6"
LIGHT = "EEF1EA"
LINE = "D9D9D9"
BLACK = "000000"
MUTED = "5E6E68"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=LINE, size="5"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=105, start=115, bottom=105, end=115):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, size=None, bold=None, color=BLACK, italic=None):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_text(doc, text, style=None, size=9.5, bold=False, color=BLACK, before=0, after=5, keep=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.12
    p.paragraph_format.keep_with_next = keep
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(8 if level == 1 else 5)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_run_font(run, size=14 if level == 1 else 11, bold=True, color=BLACK)
    return p


def style_table(table, header=True, font_size=8.3, first_col_bold=False):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for r_index, row in enumerate(table.rows):
        for c_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            if header and r_index == 0:
                set_cell_shading(cell, FOREST)
            elif r_index % 2 == 0:
                set_cell_shading(cell, "F7F7F3")
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                for run in p.runs:
                    set_run_font(
                        run,
                        size=font_size,
                        bold=(header and r_index == 0) or (first_col_bold and c_index == 0),
                        color="FFFFFF" if header and r_index == 0 else BLACK,
                    )


def add_bullet(doc, lead, body):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.line_spacing = 1.05
    r1 = p.add_run(lead + ": ")
    set_run_font(r1, 9, True)
    r2 = p.add_run(body)
    set_run_font(r2, 9, False)
    return p


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.top_margin = Inches(0.55)
section.bottom_margin = Inches(0.55)
section.left_margin = Inches(0.66)
section.right_margin = Inches(0.66)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
styles["Normal"].font.size = Pt(9.5)
for name in ("Title", "Heading 1", "Heading 2"):
    styles[name].font.name = "Arial"
    styles[name]._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    styles[name]._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    styles[name].font.color.rgb = RGBColor(0, 0, 0)

title = doc.add_paragraph(style="Title")
title.paragraph_format.space_after = Pt(3)
run = title.add_run("Sıfırİz Marka Stratejisi")
set_run_font(run, 26, True, BLACK)

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(9)
r = subtitle.add_run("Atıksız şehir yaşamı içerik markası  |  Final proje özeti")
set_run_font(r, 10, True, MUTED)

add_text(
    doc,
    "Sıfırİz, şehirde yaşayan kişilere atık azaltma ve yeniden kullanım alışkanlıkları kazandıran kurgusal bir dijital içerik markasıdır. Marka, büyük yaşam tarzı vaatleri yerine aynı gün denenebilecek küçük davranışlar önerir. Temel sonuç şudur: sürdürülebilir seçim, kullanıcıya suçluluk yüklediğinde değil; doğru anda açık ve erişilebilir hale geldiğinde tekrarlanır.",
    size=10,
    after=7,
)

add_heading(doc, "Marka Konumlandırması", 1)
add_text(doc, "Sıfırİz; kurumsal atık yönetimi bilgisini kopyalamaz. Bu bilgiyi ev, mahalle ve paylaşım döngüsünde uygulanabilir mikro görevlere çevirir. Marka bir belediye hizmeti veya ürün pazaryeri değil, davranış odaklı bir içerik rehberidir.", size=9.3)

framework = doc.add_table(rows=1, cols=2)
framework.columns[0].width = Inches(1.28)
framework.columns[1].width = Inches(5.55)
framework.rows[0].cells[0].text = "Çerçeve"
framework.rows[0].cells[1].text = "Tanım"
rows = [
    ("Neden", "Şehir yaşamında sürdürülebilir seçimi anlaşılır ve tekrarlanabilir hale getirmek."),
    ("Kime", "22–38 yaşında, kentte yaşayan, çevreye duyarlı fakat zaman ve bilgi yükü hisseden dijital kullanıcılara."),
    ("Ne", "Kısa rehberler, haftalık mikro görevler, kaynaklı açıklamalar ve paylaşılabilir örnekler."),
    ("Nasıl", "Yargılamayan dil, tek eylemli çağrılar, gerçek ev/mahalle sahneleri ve insan onaylı içerik akışı."),
]
for label, value in rows:
    cells = framework.add_row().cells
    cells[0].text = label
    cells[1].text = value
style_table(framework, first_col_bold=True)

add_heading(doc, "Kimlik ve Görsel Dil", 1)
palette = doc.add_table(rows=2, cols=5)
labels = [("Orman", FOREST), ("Adaçayı", SAGE), ("Kayısı", APRICOT), ("Kâğıt", CREAM), ("Kömür", "20332E")]
for idx, (name, hex_value) in enumerate(labels):
    palette.rows[0].cells[idx].text = name
    palette.rows[1].cells[idx].text = f"#{hex_value}"
    set_cell_shading(palette.rows[0].cells[idx], hex_value)
    set_cell_shading(palette.rows[1].cells[idx], "FFFFFF")
    for cell in (palette.rows[0].cells[idx], palette.rows[1].cells[idx]):
        set_cell_border(cell)
        set_cell_margins(cell, 90, 80, 90, 80)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for rr in cell.paragraphs[0].runs:
            set_run_font(rr, 8.2, True, "FFFFFF" if hex_value in (FOREST, "20332E") and cell is palette.rows[0].cells[idx] else BLACK)

add_text(doc, "Görsel dil; doğal ışık, gündelik nesneler, geri dönüştürülmüş kâğıt dokuları, kusurlu gerçek malzemeler ve döngüsel kompozisyonlardan oluşur. Yaprak logosu klişesi, steril stok fotoğrafı, korku görseli ve doğrulanmamış çevre rozeti kullanılmaz.", size=9.2, before=4)

add_heading(doc, "Marka Değerleri", 1)
for lead, body in [
    ("Uygulanabilirlik", "Her içerik tek ve net bir davranış önerir"),
    ("Şeffaflık", "Kaynağı olmayan istatistik kullanılmaz; sınırlar açıklanır"),
    ("Kapsayıcılık", "Sürdürülebilirlik pahalı alışverişle eşitlenmez"),
    ("Yerellik", "Ev, mahalle, pazar, tamirci ve paylaşım noktaları görünür kılınır"),
    ("İyimser gerçekçilik", "Sorun küçümsenmez; kullanıcı suçlanmaz"),
]:
    add_bullet(doc, lead, body)

doc.add_page_break()

add_heading(doc, "Marka Sesi ve Hedef Kitle", 1)
add_text(doc, "Sıfırİz bilgili fakat öğretmen gibi konuşmayan; pratik, sakin ve cesaretlendirici bir komşu gibi yazar. Hedef persona Derya, 28 yaşında, yoğun çalışan ve içerikleri çoğunlukla telefonundan takip eden bir şehir kullanıcısıdır. Çevreye duyarlıdır ancak uzun raporlar ve kusursuz yaşam örnekleri yerine iki dakikada anlayıp aynı gün deneyebileceği önerilere ihtiyaç duyar.", size=9.3)

voice = doc.add_table(rows=1, cols=2)
voice.columns[0].width = Inches(3.42)
voice.columns[1].width = Inches(3.42)
voice.rows[0].cells[0].text = "Kullan"
voice.rows[0].cells[1].text = "Kaçın"
voice.add_row().cells[0].text = "“Bugün şunu deneyebilirsin.”\nSomut eylem ve gündelik Türkçe\nKaynağı ve kapsamı açık iddia"
voice.rows[1].cells[1].text = "“Bunu yapmalısın.”\nSuçluluk, korku ve kesin sonuç\nKusursuzluk ve mucize vaatleri"
style_table(voice, font_size=8.6)

add_heading(doc, "İçerik Sütunları", 1)
pillars = doc.add_table(rows=1, cols=2)
pillars.columns[0].width = Inches(1.75)
pillars.columns[1].width = Inches(5.1)
pillars.rows[0].cells[0].text = "Sütun"
pillars.rows[0].cells[1].text = "Kapsam"
for label, value in [
    ("Evde 10 Dakika", "Buzdolabı, yeniden doldurma, tek kullanımlık ürün ve ev içi mikro görevler."),
    ("Mahallede Döngü", "Tamir, ikinci el, paylaşım dolabı, üretici pazarı ve yerel dönüşüm noktaları."),
    ("Etikete Değil Alışkanlığa", "Yeşil iddiaları sade dille sorgulayan ve kaynağa bağlayan açıklamalar."),
    ("Paylaş ve Çoğalt", "Tamamlanan görevler, önce–sonra örnekleri ve topluluk çağrıları."),
]:
    c = pillars.add_row().cells
    c[0].text = label
    c[1].text = value
style_table(pillars, first_col_bold=True)

add_heading(doc, "Benzer Oluşumlardan Ayrışma", 1)
competitors = doc.add_table(rows=1, cols=3)
competitors.columns[0].width = Inches(1.22)
competitors.columns[1].width = Inches(2.6)
competitors.columns[2].width = Inches(3.03)
competitors.rows[0].cells[0].text = "Oluşum"
competitors.rows[0].cells[1].text = "Güçlü yön"
competitors.rows[0].cells[2].text = "Sıfırİz ayrımı"
for org, strong, diff in [
    ("Sıfır Atık", "Kurumsal ve kapsamlı atık bilgisi", "Bilgiyi gündelik mikro davranışa çevirir"),
    ("WWF Türkiye", "Araştırma ve savunuculuk", "Ertesi gün tekrarlanabilir şehir rutinlerine odaklanır"),
    ("Fazla", "Gıda fazlasını teknolojiyle yönetir", "Tek hizmet yerine ev ve mahalle davranışlarını kapsar"),
    ("Too Good To Go", "Net ürün deneyimi ve topluluk", "Tek işlem yerine davranış repertuvarı geliştirir"),
    ("UNEP Zero Waste", "Küresel veri ve politika çerçevesi", "Büyük resmi yerel Türkçe ve düşük eşikli görevlerle tamamlar"),
]:
    c = competitors.add_row().cells
    c[0].text = org
    c[1].text = strong
    c[2].text = diff
style_table(competitors, font_size=7.7, first_col_bold=True)

add_heading(doc, "Kaynaklar", 2)
sources = [
    "Sıfır Atık Türkiye — sifiratik.gov.tr/sifir-atik/sifir-atik-nedir",
    "WWF-Türkiye — wwf.org.tr/?15760/sahilsahiplenprogramsonuclar2023-2024=",
    "Fazla — fazla.com/gida/",
    "Too Good To Go — toogoodtogo.com/en-us/about-us",
    "UNEP — unep.org/interactives/zero-waste-101/",
]
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(0)
p.paragraph_format.line_spacing = 1.0
for index, source in enumerate(sources):
    run = p.add_run(("  •  " if index else "") + source)
    set_run_font(run, 7.2, False, MUTED)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = footer.add_run("Sıfırİz  ·  Yapay Zeka ile İçerik Üretimi Final Projesi  ·  2026")
set_run_font(fr, 7.5, False, MUTED)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
